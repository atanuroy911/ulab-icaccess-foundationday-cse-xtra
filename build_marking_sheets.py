"""Printable marking sheets (one page per session, A4 landscape) that mirror the web-app form."""
import html, re
from pathlib import Path
import openpyxl
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).parent
XLSX = max(ROOT.glob("ICACCESS2026-Plan_Dynamic_V4*.xlsx"), key=lambda f: f.stat().st_mtime)
wb = openpyxl.load_workbook(XLSX, data_only=True)
cl = lambda x: re.sub(r"\s+", " ", str(x or "")).strip()
NAME_FIXES = {"Dr. Mohammaed Ashikur Rahman": "Dr. Mohammed Ashikur Rahman", "Dr. Md.Ruhul Amin": "Dr. Md. Ruhul Amin"}
fix = lambda n: NAME_FIXES.get(n, n)
code_gs = (ROOT / "chair-eval-webapp" / "Code.gs").read_text(encoding="utf-8")
grab = lambda name: re.findall(r"'([^']+)'", re.search(rf"const {name} = \[(.*?)\];", code_gs, re.S).group(1))
CRIT, HELP, SCALE = grab("CRITERIA"), grab("CRITERIA_HELP"), grab("SCALE")

titles = {r[1]: cl(r[2]) for r in wb["Sessions"].iter_rows(min_row=2, values_only=True) if r[1]}
tracks = {r[1]: cl(r[0]) for r in wb["Sessions"].iter_rows(min_row=2, values_only=True) if r[1]}
codes = [c.value.strip() for c in wb["_Lists"]["A"][1:] if c.value and str(c.value).strip()]
sessions = []
for code in codes:
    ws = wb[code]
    first = next((r + 1 for r in range(1, 40) if cl(ws.cell(r, 1).value) == "ID"), 12)
    head_end = next((r for r in range(1, 40) if cl(ws.cell(r, 1).value) in ("ID", "Papers")), 11)
    kv = {cl(ws.cell(r, 1).value): ws.cell(r, 2).value for r in range(1, head_end)}
    chairs = [fix(cl(ws.cell(r, 2).value).replace("Md.Ruhul", "Md. Ruhul")) for r in range(1, head_end)
              if cl(ws.cell(r, 1).value).startswith("Session Chair") and cl(ws.cell(r, 2).value)]
    papers = [(int(ws.cell(r, 1).value), cl(ws.cell(r, 2).value)) for r in range(first, ws.max_row + 1)
              if str(ws.cell(r, 1).value or "").strip().isdigit() and ws.cell(r, 2).value]
    if papers:
        sessions.append(dict(code=code, title=cl(kv.get("Session Title")) or titles.get(code, ""), track=tracks.get(code, ""),
                             chairs=chairs, time=cl(kv.get("Session Time")), room=cl(kv.get("Session Room")), papers=papers))

e = html.escape
box = '<span class="box"></span>'
n = len(CRIT)
head_cells = "".join(f'<th class="q"><b>{i + 1}</b> {e(c)}<small>{e(h)}</small></th>' for i, (c, h) in enumerate(zip(CRIT, HELP)))
scale = "".join(f'<span><b>{i}</b> {e(s)}</span>' for i, s in enumerate(SCALE))


def table(rows, manual=False):
    rh = min(28, 122 // max(len(rows), 1))
    body = ""
    for pid, title in rows:
        body += (f'<tr style="height:{rh}mm"><td class="pid">{pid}</td><td class="ttl">{e(title)}</td>'
                 + "".join(f'<td class="sc">{box}</td>' for _ in CRIT)
                 + f'<td class="tot"><span class="box wide"></span><small>/ {5 * n}</small></td>'
                 + '<td class="vote"><span>Yes</span><span>No</span></td><td class="cm"></td></tr>')
    return (f'<table><thead><tr><th class="pid">ID</th><th class="ttl">Paper title</th>{head_cells}'
            f'<th class="tot">Total</th><th class="vote">Best paper vote</th><th class="cm">Session chair comments</th></tr></thead>'
            f'<tbody>{body}</tbody></table>')


def page(meta, rows, cont=False, extra="", manual=False):
    where = "Online" if meta["room"].lower().startswith("http") else meta["room"]
    return f'''<section class="page{' dense' if len(rows) >= 7 else ''}">
  <header>
    <img src="icaccess-logo.png" alt="">
    <div class="h1"><p class="eyebrow">iCACCESS 2026 · October 2–3, 2026 · ULAB Campus, Dhaka</p>
      <h1>Session Chair Marking Sheet{extra}{" <em>(continued)</em>" if cont else ""}</h1></div>
    <div class="scale"><p>Score each criterion from 0 to 5</p><div>{scale}</div></div>
  </header>
  {MANUAL_META if manual else ""}
  <div class="meta" {'hidden' if manual else ""}>
    <div><b>Session</b> {e(meta["code"])}{(" · " + e(meta["title"])) if meta["title"] else ""}</div>
    <div><b>Track</b> {e(meta["track"]) or "&nbsp;"}</div>
    <div><b>Time / room</b> {e(meta["time"])} {("· " + e(where)) if where else ""}</div>
    <div><b>Scheduled chairs</b> {e(" & ".join(meta["chairs"])) or "&nbsp;"}</div>
    <div class="line"><b>Chair name (if different)</b></div>
  </div>
  {table(rows, manual)}
  <footer><span>Total = sum of the {n} scores (maximum {5 * n}). Skip papers that were not presented.</span>
    <span class="sig"><b>Signature</b></span></footer>
</section>'''


MANUAL_META = ('<div class="meta manual"><div class="f"><b>Session ID</b></div><div class="f"><b>Track</b></div>'
               '<div class="f"><b>Time / room</b></div><div class="f w2"><b>Session chair name(s)</b></div>'
               '<div class="f"><b>Signature</b></div></div>')
pages = [page(s, s["papers"]) for s in sessions]          # one page per session, however many papers
blank = dict(code="", title="", track="", chairs=[], time="", room="")
for _ in range(2):                                         # manual sheets: everything written by hand
    pages.append(page(blank, [("", "")] * 7, extra=" · manual", manual=True))

css = '''
@page { size: A4 landscape; margin: 0 }
* { box-sizing: border-box; margin: 0 }
body { font-family: 'Atkinson Hyperlegible', Arial, sans-serif; color: #0f172a; -webkit-print-color-adjust: exact; print-color-adjust: exact }
.page { width: 297mm; height: 210mm; padding: 9mm 10mm 8mm; display: flex; flex-direction: column; page-break-after: always; overflow: hidden }
header { display: flex; align-items: center; gap: 6mm; border-bottom: 0.6mm solid #0B1F3A; padding-bottom: 2.5mm }
header img { height: 14mm }
.h1 { flex: 1 }
.eyebrow { font-size: 8pt; color: #0A6B60; font-weight: 700 }
h1 { font-family: Outfit, sans-serif; font-size: 17pt; color: #0B1F3A; line-height: 1.1 }
h1 em { font-size: 11pt; font-weight: 500; color: #64748b }
.scale { text-align: right; font-size: 7.5pt }
.scale p { font-weight: 700; color: #0B1F3A; margin-bottom: 1mm }
.scale div { display: flex; gap: 2.5mm; justify-content: flex-end; flex-wrap: wrap }
.scale b { display: inline-grid; place-items: center; width: 4mm; height: 4mm; border-radius: 1mm; background: #0B1F3A; color: #fff; font-size: 7pt }
.meta { display: grid; grid-template-columns: 1.5fr 1.2fr 1fr; gap: 1.5mm 6mm; margin: 3mm 0; font-size: 9pt }
.meta b { color: #475569; font-size: 7.5pt; text-transform: uppercase; letter-spacing: .06em; margin-right: 1.5mm }
.meta .line { grid-column: 3; grid-row: 2; border-bottom: 0.3mm solid #475569; align-self: end }
table { width: 100%; border-collapse: collapse; table-layout: fixed }
.manual { display: grid; grid-template-columns: 1fr 1fr 1fr 1.6fr 1fr; gap: 0 6mm; margin: 4mm 0 3mm; font-size: 9pt }
.manual .f { border-bottom: 0.3mm solid #475569; height: 8mm; display: flex; align-items: end; padding-bottom: 0.5mm }
.manual .w2 { grid-column: span 1 }
[hidden] { display: none !important }
.dense { padding-top: 6mm; padding-bottom: 5mm }
.dense header img { height: 11mm }
.dense .meta { margin: 2mm 0 }
.dense td { padding: 0.8mm 1.2mm }
.dense td.ttl { font-size: 7pt; line-height: 1.1 }
.dense .box { width: 8.5mm; height: 8.5mm }
.dense td.vote span { margin: 0.6mm 0 }
.dense td.tot small { display: none }
.dense footer { margin-top: 1.5mm }
th, td { border: 0.3mm solid #94a3b8; padding: 1.5mm; vertical-align: middle }
thead th { background: #EEF2F8; color: #0B1F3A; font-size: 7.5pt; text-align: center; font-weight: 700; line-height: 1.15 }
thead th small { display: block; font-weight: 400; color: #475569; font-size: 5.8pt; margin-top: 0.8mm; line-height: 1.15 }
th.pid, td.pid { width: 13mm; text-align: center; font-family: Outfit, sans-serif; font-weight: 700; font-size: 12pt; color: #1E5AA8 }
th.ttl, td.ttl { width: 62mm; font-size: 8pt; line-height: 1.2 }
th.q, td.sc { width: 17.5mm }
th.tot, td.tot { width: 19mm; text-align: center }
th.vote, td.vote { width: 19mm }
td.sc, td.tot { text-align: center }
.box { display: inline-block; width: 10mm; height: 10mm; border: 0.5mm solid #0B1F3A; border-radius: 1.5mm; background: #fff }
.box.wide { width: 12mm }
td.tot small { display: block; font-size: 7pt; color: #475569; margin-top: 0.5mm }
td.vote { text-align: center }
td.vote span { display: flex; align-items: center; justify-content: center; gap: 1.5mm; font-size: 8.5pt; font-weight: 700; margin: 1.5mm 0 }
td.vote span::before { content: ''; width: 4mm; height: 4mm; border: 0.5mm solid #0B1F3A; border-radius: 50% }
footer { display: flex; justify-content: space-between; align-items: end; margin-top: 3mm; font-size: 7.5pt; color: #475569 }
.sig { width: 70mm; border-bottom: 0.3mm solid #475569; text-align: left; padding-bottom: 0.5mm }
.sig b { text-transform: uppercase; letter-spacing: .06em; font-size: 7pt }
'''
doc = (f'<!doctype html><html><head><meta charset="utf-8"><title>iCACCESS 2026 – Marking Sheets</title>'
       '<link href="https://fonts.googleapis.com/css2?family=Atkinson+Hyperlegible:wght@400;700&family=Outfit:wght@500;700&display=swap" rel="stylesheet">'
       f'<style>{css}</style></head><body>{"".join(pages)}</body></html>')
out = ROOT / "ICACCESS2026_Marking_Sheets.html"
out.write_text(doc, encoding="utf-8")
pdf = ROOT / "ICACCESS2026_Marking_Sheets.pdf"
with sync_playwright() as p:
    b = p.chromium.launch(); pg = b.new_page()
    pg.goto(out.as_uri()); pg.wait_for_timeout(1500)
    pg.pdf(path=str(pdf), width="297mm", height="210mm", print_background=True, prefer_css_page_size=True)
    b.close()
print(len(sessions), "sessions,", len(pages), "pages")
